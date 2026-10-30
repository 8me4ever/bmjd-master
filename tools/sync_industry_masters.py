"""Rebase all eleven industry catalogs on the user-approved BMJD master.

Codes and all non-JD job fields come exclusively from the approved master.
Existing industry JD fields are carried forward by stable ID, never old codes.
"""
import argparse
import collections
import copy
import csv
import hashlib
import io
import json
from pathlib import Path
from split_industry_masters import partition,INDUSTRIES

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'行业子库同步_20261030'
MASTER=ROOT/'原版分类编码收口_20261030/主库.json'
OLD_MASTER=ROOT/'岗位JD匹配审查_20261029/岗位JD主库.json'
OLD=ROOT/'行业专业JD优化_20261029'
ROUTES=ROOT/'全量名称代码统一_20261029/来源去向映射.json'
VERSION='BMJD-INDUSTRY-SYNC-20261030-01'
JD_FIELDS={'jd','jd_clauses','jd_version','industry_jd_revision'}
HEADER=['岗位ID','岗位代码','岗位名称','规范职能','浏览分类','专业方向','职责系列','层级','所属行业','全部适用行业','JD类型','岗位JD','主库版本','子库版本']

def read(p):return json.loads(p.read_text('utf-8'))
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def text_sha(s):return hashlib.sha256(s.encode()).hexdigest()
def encode(d):return (json.dumps(d,ensure_ascii=False,indent=2)+'\n').encode()


def carry_forward(parent,old_child,original_parent,membership):
    if any(parent[k]!=old_child[k] or original_parent[k]!=old_child[k] for k in ['id','name','level','function_id','series_id','source_ids']):
        raise ValueError('岗位身份不一致，禁止按旧代码或近似名称接入行业JD。')
    if parent['jd']!=original_parent['jd']:
        raise ValueError('主库JD发生变化，不能静默接入旧行业修订。')
    if old_child['industry_membership']!=membership:
        raise ValueError('本行业来源范围改变，须重新核查专业JD归属。')
    child=copy.deepcopy(parent);child['industry_membership']=copy.deepcopy(membership)
    if old_child['jd']!=original_parent['jd']:
        revision=old_child.get('industry_jd_revision')
        if not revision or revision['parent_jd_sha256']!=text_sha(parent['jd']):
            raise ValueError('行业JD缺少匹配当前主库的职责版本依据。')
        for k in JD_FIELDS:
            if k in old_child:child[k]=copy.deepcopy(old_child[k])
    else:
        if any(old_child[k]!=original_parent[k] for k in ['jd','jd_clauses','jd_version']):
            raise ValueError('未改写岗位的JD字段与原父库不一致。')
    return child


def table(child):
    return [HEADER]+[[j['id'],j['code'],j['name'],j['function_name'],j['function_group_name'],j['function_direction_name'],j['series_name'],str(j['level']),child['industry'],'、'.join(j['industries']),'保留行业专业JD' if 'industry_jd_revision' in j else '沿用主库JD',j['jd'],child['parent_version'],child['version']] for j in child['jobs']]


def tsv(rows):
    s=io.StringIO(newline='');csv.writer(s,delimiter='\t',lineterminator='\n').writerows(rows);return s.getvalue().encode('utf-8-sig')


def build():
    master=read(MASTER);original=read(OLD_MASTER);old_by={j['id']:j for j in original['jobs']}
    routes=read(ROUTES);parts=partition(master['jobs'],routes);parent_sha=sha(MASTER)
    files={};summaries=[];migration=[];specialized=[];inputs=[MASTER,OLD_MASTER,ROUTES,OLD/'子库优化总览.json',Path(__file__)]
    for industry in INDUSTRIES:
        source=OLD/industry/'岗位主库.json';inputs.append(source);old=read(source);old_jobs={j['id']:j for j in old['jobs']}
        jobs=[]
        assert [j['id'] for j in parts[industry]]==[j['id'] for j in old['jobs']]
        for parent in parts[industry]:
            clean={k:v for k,v in parent.items() if k!='industry_membership'}
            child=carry_forward(clean,old_jobs[parent['id']],old_by[parent['id']],parent['industry_membership']);jobs.append(child)
            changed=child['jd']!=clean['jd'];previous=old_jobs[child['id']]
            migration.append({'industry':industry,'job_id':child['id'],'name':child['name'],'previous_code':previous['code'],'current_code':child['code'],
              'code_changed':previous['code']!=child['code'],'industry_jd_retained':changed,'previous_child_version':old['version'],'child_version':VERSION,
              'parent_version':master['version'],'jd_sha256':text_sha(child['jd'])})
            if changed:specialized.append({'industry':industry,'job_id':child['id'],'name':child['name'],'code':child['code'],'jd_version':child['jd_version'],
              'parent_jd_sha256':text_sha(clean['jd']),'industry_jd_sha256':text_sha(child['jd']),'source_child_path':source.relative_to(ROOT).as_posix(),
              'source_child_sha256':sha(source),'source_code':previous['code'],'local_source_ids':child['industry_membership']['source_ids']})
        ids={j['id'] for j in jobs};functions={}
        for j in jobs:functions.setdefault(j['function_id'],{'id':j['function_id'],'name':j['function_name'],'job_ids':[]})['job_ids'].append(j['id'])
        series=[dict(s,job_ids=[i for i in s['job_ids'] if i in ids]) for s in master['series'] if any(i in ids for i in s['job_ids'])]
        groups=[];directions=[]
        for g in master['function_groups']:
            local=[j for j in jobs if j['function_group_id']==g['id']]
            if local:groups.append(dict(g,jobs=len(local),professional_directions=len({j['function_direction_id'] for j in local}),series=len({j['series_id'] for j in local}),code_blocks=dict(sorted(collections.Counter(j['code'][:2] for j in local).items()))))
        for d in master['function_directions']:
            local=[j for j in jobs if j['function_direction_id']==d['id']]
            if local:directions.append(dict(d,jobs=len(local),series=len({j['series_id'] for j in local})))
        local_routes=[r for r in routes if r['industry']==industry]
        count=sum(j['jd']!=old_by[j['id']]['jd'] for j in jobs)
        summary={'industry':industry,'jobs':len(jobs),'unique_codes':len({j['code'] for j in jobs}),'functions':len(functions),'function_groups':len(groups),
          'professional_directions':len(directions),'series':len(series),'original_source_records':len(local_routes),
          'active_original_sources':sum(r['state']!='用户暂缓' for r in local_routes),'paused_original_sources':sum(r['state']=='用户暂缓' for r in local_routes),
          'industry_jd_retained':count,'master_jd_retained':len(jobs)-count,'path':f'{industry}/岗位主库.json'}
        child={'version':VERSION,'name_code_version':master['name_code_version'],'tier':1,'industry':industry,'parent_version':master['version'],
          'parent_master_path':MASTER.relative_to(ROOT).as_posix(),'parent_master_sha256':parent_sha,'source_industry_version':old['version'],
          'source_industry_path':source.relative_to(ROOT).as_posix(),'source_industry_sha256':sha(source),
          'membership_basis':'原BMJD11行业活动来源映射；同岗跨行业使用同一主库身份与五位码。',
          'jd_policy':'保留已审查的行业专业JD；其他岗位沿用主库JD。本轮不新增JD改写、不回写主库。',
          'summary':summary,'jobs':jobs,'functions':list(functions.values()),'series':series,'function_groups':groups,'function_directions':directions,'sources':local_routes}
        files[f'{industry}/岗位主库.json']=encode(child);files[f'{industry}/岗位名称代码及JD.tsv']=tsv(table(child));summaries.append(summary)
    summary={'industry_subcatalogs':11,'master_jobs':9111,'industry_job_memberships':sum(r['jobs'] for r in summaries),
      'shared_master_jobs':sum(len(j['industries'])>1 for j in master['jobs']),'active_sources':19226,'paused_sources':23,
      'retained_industry_specialized_jds':len(specialized),'retained_specialized_job_ids':len({r['job_id'] for r in specialized}),
      'new_jd_rewrites':0,'master_changes':0,'child_names_or_levels_changed':0,'child_codes_from_master':True}
    assert summary['industry_job_memberships']==19018 and len(specialized)==843 and summary['retained_specialized_job_ids']==832
    overview={'version':VERSION,'parent_version':master['version'],'name_code_version':master['name_code_version'],'parent_master_sha256':parent_sha,
      'approved_master':True,'summary':summary,'industries':summaries}
    report={'version':VERSION,'all_passed':True,'summary':summary,'input_sha256':{p.relative_to(ROOT).as_posix():sha(p) for p in inputs},
      'scope':'确认主库拆分与行业专业JD迁移，不改主库、岗位身份或旧行业文件。'}
    files['子库总览.json']=encode(overview);files['核验报告.json']=encode(report);files['行业迁移对照.json']=encode(migration);files['行业JD保留索引.json']=encode(specialized)
    files['岗位行业索引.json']=encode([{'job_id':j['id'],'code':j['code'],'name':j['name'],'industries':j['industries'],'paths':[f'{i}/岗位主库.json' for i in j['industries']]} for j in master['jobs']])
    files['主库确认记录.json']=encode({'date':'2026-10-30','version':master['version'],'master_path':MASTER.relative_to(ROOT).as_posix(),'master_sha256':parent_sha,
      'basis':'用户在本会话确认该版主库，并指示以此拆分子库及同步HTML。','user_instruction':'我目前看着没什么问题了，以这个为主库拆分子库并同步到html吧','approved_scope':'11行业子库与查岗HTML；主库文件原样固定，新增工具不在本次迁移范围。'})
    return files,report


def export():
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    def sheet(wb,name,rows):
        ws=wb.create_sheet(name)
        for n,row in enumerate(rows,1):
            ws.append(row)
            for c,v in enumerate(row,1):
                if isinstance(v,str):ws.cell(n,c).data_type='s'
        ws.freeze_panes='D2';ws.auto_filter.ref=ws.dimensions
        for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='254B70')
        for n in range(1,len(rows[0])+1):ws.column_dimensions[ws.cell(1,n).column_letter].width=22 if n!=12 else 95
        for row in ws.iter_rows():
            for c in row:c.alignment=Alignment(vertical='top',wrap_text=True)
    overview=read(OUT/'子库总览.json');combined=Workbook();combined.remove(combined.active)
    stats=[['行业','岗位数','行业专业JD保留','沿用主库JD','主库版本','子库版本']]+[[r['industry'],r['jobs'],r['industry_jd_retained'],r['master_jd_retained'],overview['parent_version'],VERSION] for r in overview['industries']]
    sheet(combined,'11行业总览',stats)
    for i in INDUSTRIES:
        child=read(OUT/i/'岗位主库.json');rows=table(child);sheet(combined,i,rows)
        wb=Workbook();wb.remove(wb.active);sheet(wb,'岗位主库',rows);sheet(wb,'版本与统计',stats[:1]+[r for r in stats[1:] if r[0]==i]);wb.save(OUT/i/'岗位主库.xlsx')
    combined.save(OUT/'BMJD_11行业子库.xlsx')


def main():
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--check',action='store_true');a=p.parse_args();files,report=build()
    if a.check:
        for n,b in files.items():assert (OUT/n).read_bytes()==b,n
    else:
        OUT.mkdir(exist_ok=True)
        for n,b in files.items():
            path=OUT/n
            if path.exists():assert path.read_bytes()==b,'不覆盖已有不同版本：'+n
            else:path.parent.mkdir(parents=True,exist_ok=True);path.write_bytes(b)
        export()
    print(json.dumps(report['summary'],ensure_ascii=False))


if __name__=='__main__':main()
