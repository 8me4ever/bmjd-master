"""Review every 3–6 level slot and build an additive, traceable BMJD master v2.

No original role is renamed, recoded, moved or rewritten. Added roles are explicit
professional inferences; legacy source records are context, not invented evidence.
"""
import argparse
import collections
import copy
import csv
import hashlib
import io
import json
import re
import unicodedata
from pathlib import Path

ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'主库v2_层级补全_20261030'
BASE=ROOT/'原版分类编码收口_20261030/主库.json';BACKUP=ROOT/'主库备份_20261030_v1';POLICY=OUT/'补全规则.json'
JD_RULES=ROOT/'岗位JD匹配审查_20261029/JD审查与编写规则.json'
VERSION='BMJD-MASTER-V2-LEVELS-20261030-01'
QUAL=re.compile(r'[（(][^（）()]*[1-7]级[^（）()]*[）)]')
GRADES=re.compile(r'^(?:资深|高级|中级|初级|助理|首席|中高级|初中级)+')
OCCUPATION=re.compile(r'(?:工程师|设计师|建筑师|规划师|教师|讲师|教练|顾问|研究员|分析师|审计师|会计师|精算师|律师|咨询师|经济师|评估师|测量师|技师|检验员|化验员|试验员|测绘员|检测员|监理员|管理员|编辑|记者|翻译|编审|导演|编剧|摄影师|摄像师|灯光师|美术师|化妆师|演员|歌手|乐师|调音师|厨师|茶艺师|品酒师|咖啡师|营养师|美容师|美发师|按摩师|理疗师|治疗师|康复师|业务员|销售员|代表|柜员|交易员|经纪人|统计员|美工|师)$')
FIXED=re.compile(r'(?:总经理|总裁|首席.{0,12}官|董事(?:长|会秘书)?|监事(?:会主席)?|校长|院长|行长|书记|秘书长|总工程师|总建筑师|总设计师)')
SKILLED=re.compile(r'(?:电工|钳工|焊工|铆工|车工|铣工|磨工|木工|模具工|机修工|维修工|制版工|裁剪工|缝纫工|雕刻工|铸造工|锻造工|轧钢工|炼钢工|炼铁工)$')
OPERATOR=re.compile(r'(?:工|操作员|主操|司机|服务员|清洁员|保洁员|保安员|值班员|门卫|迎宾员|传菜员|收银员|保管员)$')
CLERICAL=re.compile(r'(?:文员|秘书|接线员|话务员|助理|见习生|实习生|学徒)$')
PROFILE_SUBSTITUTIONS={'护理管理':'护理'}

def read(p):return json.loads(p.read_text('utf-8'))
def encode(d):return (json.dumps(d,ensure_ascii=False,indent=2)+'\n').encode()
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def norm(s):return ''.join(c for c in unicodedata.normalize('NFKC',s) if not c.isspace() and unicodedata.category(c)!='Cf').casefold()
def split_name(name):
    value=QUAL.sub('',name).strip()
    if '/' in value and re.search(r'(?:师|员|经理|主管|总监|助理|主任)$',value.split('/')[0]):value=value.split('/')[0]
    match=re.search(r'([（(][^（）()]+[）)])$',value)
    return (value[:match.start()].strip(),match[1]) if match else (value,'')
def ungrade(name):return GRADES.sub('',name).replace('助理工程师','工程师')
def scope_for(s):
    core,context=split_name(s['name']);core=ungrade(core)
    core=re.sub(r'(?:副主任医师|主任医师|主治医师|住院医师|主任医生|医生|医师|副主任技师|主任技师|主治技师|主管技师|技师|技士|副主任|主任|主治)$','',core)
    if core!='美工':core=re.sub(r'(?:总监|经理|主管|专员|助理|工程师|设计师|管理员|操作工|操作员|班长|组长|领班|队长|主操|工)$','',core)
    return core or s['function_name'],context


def route(s,jobs,policy,issue_ids):
    scope,context=scope_for(s);cores=[split_name(j['name'])[0] for j in jobs]
    anchors=[{'job_id':j['id'],'level':j['level'],'name':j['name'],'jd_profile':j['jd_profile']} for j in jobs]
    result={'type':'management','scope':scope,'context':context,'base':'','anchors':anchors,'reason':'沿同系列职能管理梯队补齐。'}
    if s['function_name'] in policy['fixed_governance_functions'] or FIXED.search(s['name']) and not any(j['level']>=4 and not FIXED.search(j['name']) for j in jobs) or all(j['level']<=2 and FIXED.search(j['name']) for j in jobs):
        return dict(result,type='fixed_position',reason='治理、高管或校级等固定职务不是一般专员梯队；不制造初级总经理、助理校长等岗位。')
    if any(j['id'] in issue_ids for j in jobs):return dict(result,type='known_jd_issue',reason='本系列含已标记的旧JD错配，避免将错误职责框架扩散到新岗位。')
    if any(t in s['name']+' '.join(cores) for t in policy['paused_topics']):return dict(result,type='paused_topic',reason='属于原暂缓的叉车/搬运议题，不借补层级恢复暂缓范围。')
    physician=any(re.search(r'医师|医生|校医|法医',n) for n in cores)
    technician=any(re.search(r'技师|技士',n) for n in cores)
    if physician:
        mixed=any(re.search(r'技师|技士|技术员|卫生技术员',n) for n in cores)
        expected={2:'主任医师',3:'副主任医师',4:'主治医师',5:'住院医师'}
        consistent=all((match:=re.search(r'(副主任|主任|主治|住院)医师$',n)) is not None and {'主任':2,'副主任':3,'主治':4,'住院':5}[match[1]]==j['level'] for j,n in zip(jobs,cores))
        return dict(result,type='clinical' if consistent and not mixed else 'medical_boundary',reason='沿既有临床职称梯队审查；住院医师已为入门岗位，不为填6级另造医师助理。' if consistent and not mixed else '医师与技术员混合，或既有职称/层级为特殊口径；需先核实专业梯队，不按通用职衔填空。')
    if any(j['jd_profile']=='药师' for j in jobs):
        if any('驻店' in n or '执业' in n for n in cores):return dict(result,type='medical_boundary',reason='驻店/执业药师为特定执业岗位，不套出不具同等执业职责的低层药士或经理。')
        raw=next(n for n in cores if '药师' in n or '药剂师' in n);raw=re.sub(r'[-—](?:初级|中级|高级)$','',raw)
        match=re.search(r'(副主任|主任|主治|主管)?(中药师|西药师|药剂师|药师)$',raw)
        if match:return dict(result,type='pharmacy',scope=raw[:match.start()],base='药师' if match[2]=='药剂师' else match[2],reason='药学沿副主任/主管药师、药师与药士梯队；主管药师是专业职称，不改成经理。')
    if technician and any(j['jd_profile'] in {'医学检验','医学影像','康复'} for j in jobs):
        if not re.search(r'检验|影像|放射|CT|核磁|MR|X线|超声|输血|康复|理疗|针灸|按摩|病理技术|电生理|心电|脑电',scope,re.I):
            return dict(result,type='medical_boundary',reason='既有医疗技师职称所在专业需先核实；不把临床科室名称直接推为通用技士岗位。')
        return dict(result,type='medical_tech',reason='沿原医技专业梯队补副主任/主管技师、技师及技士，保留原有职称与代码。')
    if any(re.search(r'护士|护师',n) for n in cores):return dict(result,type='nursing',reason='护理保留专业准入与辅助边界，不把科护士长机械扩成经营经理。')
    if all(n.endswith('助理') for n in cores):
        raw,c=split_name(min(jobs,key=lambda j:abs(j['level']-6))['name']);base=ungrade(raw[:-2])
        if OCCUPATION.search(base):return dict(result,type='professional',base=base,context=c,reason='此行只保留了专业助理，按原专业本岗梯队查缺；不制造高级助理。')
        return dict(result,scope=base,context=c,reason='此行仅有职能助理，回到同职能经理/主管/专员检查，不派生高级助理。')
    if any('司机' in n for n in cores):
        raw,c=split_name(next(j['name'] for j in jobs if '司机' in j['name']));base=re.sub(r'司机(?:班长|队长|车队长)$','司机',ungrade(raw));scope=re.sub(r'司机.*$','',base) or '车辆'
        return dict(result,type='driving',base=base,scope=scope,context=c,reason='驾驶与车队基层管理分层，不设置3级司机经理；沿实际车型去掉系列旧标签中的重复车字。')
    skilled=next((j for j in sorted(jobs,key=lambda j:abs(j['level']-6)) if SKILLED.search(ungrade(split_name(j['name'])[0]))),None)
    if skilled:
        core,c=split_name(skilled['name']);base=ungrade(core)
        return dict(result,type='skilled_trade',base=base,context=c,scope=base,reason='技术工种按高级技师/技师及高级工/中级工梯队补充，不将工人改称工程师。')
    # Prefer an actual professional title near level 5, not a truncated series
    # label such as 企业内刊主编 or a management title from the top of the row.
    professional=next((j for j in sorted(jobs,key=lambda j:(abs(j['level']-5),j['level'])) if OCCUPATION.search(ungrade(split_name(j['name'])[0]))),None)
    if professional:
        core,c=split_name(professional['name']);base=ungrade(core)
        if base.endswith('厨师'):
            base=re.sub(r'(?:资深|高级|中级|初级)(?=厨师$)','',base)
            return dict(result,type='cooking',base=base,scope=base[:-2],context=c,reason='厨师沿总厨、厨师长、厨师与初级厨师称谓，不造厨师经理。')
        shifted=any((j['level']==3 and split_name(j['name'])[0].startswith('高级') or j['level']==4 and split_name(j['name'])[0]==base) and ungrade(split_name(j['name'])[0])==base for j in jobs)
        management3=not OCCUPATION.search(scope) and any(j['level']<=2 and re.search(r'总监|部长|负责人',j['name']) or j['level']==4 and re.search(r'主管|主任',j['name']) and not re.search(r'医师|技师|药师|护师',j['name']) for j in jobs)
        return dict(result,type='shifted_professional' if shifted else 'professional',base=base,context=c,management3=management3,
                    reason='按本系列现有专业称谓及层级锚点延伸；保留专业路线，不统一改成专员。')
    operator=next((j for j in sorted(jobs,key=lambda j:(-j['level'])) if OPERATOR.search(ungrade(split_name(j['name'])[0]))),None)
    if operator:
        raw,c=split_name(operator['name']);return dict(result,type='operations',base=ungrade(raw),context=c,reason='基础作业/服务方向补基层主管、组长或熟练作业岗，不机械设置3级经营岗位。')
    clerical=next((j for j in jobs if CLERICAL.search(split_name(j['name'])[0]) and j['level']>=5),None)
    if clerical and all(CLERICAL.search(n) for n in cores) and not any(j['level']<=4 for j in jobs):
        core,c=split_name(clerical['name']);return dict(result,type='clerical',base=ungrade(core),context=c,reason='专门文职或辅助岗位按高级/本岗/助理延伸，不设3级经营职务。')
    if any(j['level']>=5 and re.search(r'经理|总监',split_name(j['name'])[0]) for j in jobs):return dict(result,type='title_level_boundary',reason='岗位在5/6级却使用经理/总监职衔；先核实其专业与管理边界，不造重名异级岗位。')
    return result


def candidate(route,level,policy):
    t=route['type'];scope=route['scope'];context=route['context'];base=route['base']
    if t in {'fixed_position','known_jd_issue','paused_topic','medical_boundary','title_level_boundary'}:return None,route['reason']
    if t=='clinical':
        if level==6:return None,'临床医生现有5级住院医师为入门位置；不虚设6级独立医师或医师助理。'
        return scope+policy['clinical_levels'][str(level)]+context,route['reason']
    if t=='medical_tech':return policy['grade_patterns'][t][str(level)].format(scope=scope)+context,route['reason']
    if t=='pharmacy':return scope+({3:'副主任'+base,4:'主管'+base,5:base,6:base[:-1]+'士'}[level])+context,route['reason']
    if t=='cooking':return scope+({3:'行政总厨',4:'厨师长',5:'厨师',6:'初级厨师'}[level])+context,route['reason']
    if t=='nursing':
        scope=re.sub(r'护士长|护理部|护士|护师','',scope) or '护理'
        if level==3:return None,'现有科护士长为4级，缺乏独立更高管理职务依据；保留该专业边界，不额外造护理经理。'
        return (scope+'护士长' if level==4 else scope+'护士' if level==5 else scope+'助理' if scope.endswith('护理') else scope+'护理助理')+context,route['reason']
    if t=='skilled_trade':return ({3:base+'高级技师',4:base+'技师',5:'高级'+base,6:'中级'+base}[level])+context,route['reason']
    if t=='operations':
        if level==3:return None,'基础作业/服务系列没有独立经营层职责依据，3级保留空位；基层管理与技能执行另补。'
        action='' if scope.endswith(('操作','作业','服务')) else '作业'
        return ({4:scope+action+'主管',5:scope+action+'组长',6:'熟练'+base}[level])+context,route['reason']
    if t=='driving':
        if level==3:return None,'司机为驾驶执行岗位，不额外造3级司机经理；车队管理已有独立专业方向。'
        return ({4:scope+'车队主管',5:scope+'驾驶班组长',6:'熟练'+base}[level])+context,route['reason']
    if t=='clerical':
        if level==3:return None,'单一文职/辅助岗位不机械设置3级经营或资深专家职务。'
        name={4:'高级'+base,5:base,6:base if base.endswith('助理') else base+'助理'}[level]
        return name+context,route['reason']
    if t in {'professional','shifted_professional'}:
        if level==3 and route.get('management3') and t!='shifted_professional':return scope+'经理'+context,'本系列已有总监或主管管理锚点，为缺失3级补职能经理。'
        return policy['grade_patterns'][t][str(level)].format(base=base)+context,route['reason']
    return policy['grade_patterns']['management'][str(level)].format(scope=scope)+context,route['reason']


def role_mode(name,level,kind):
    if kind in {'clinical','medical_tech','pharmacy'}:return '资深专业' if level in {3,4} else '辅助或初级' if level==6 else '专业执行'
    if kind in {'operations','driving'} and level==6:return '技能执行'
    if level==6:return '辅助或初级'
    if re.search(r'经理|主管|护士长|负责人|组长',name):return '管理'
    return '资深专业' if level in {3,4} else '专业执行'


def make_jd(name,level,r,jobs,policy,jd_rules):
    near=min(jobs,key=lambda j:(abs(j['level']-level),j['level']))
    profile=near['jd_profile'];profile=PROFILE_SUBSTITUTIONS.get(profile,profile)
    mode=role_mode(name,level,r['type']);scope=r['scope'];custom=r['type']=='nursing' and level==6
    if r['type']=='nursing' and not custom:profile='护理'
    if r['type']=='driving':profile='驾驶' if level==6 else '运输调度'
    specific=policy.get('series_jd_overrides',{}).get(jobs[0]['series_id'])
    if custom:texts=list(policy['medical_assistance_jd']);profile='v2护理辅助'
    elif r['type']=='pharmacy' and level==6:texts=list(policy['pharmacy_assistance_jd']);profile='v2药学辅助';custom=True
    else:
        if specific:
            profile=specific['profile'];source_texts=specific.get('clauses',jd_rules['profiles'].get(profile,{}).get('clauses',[]))
        else:source_texts=jd_rules['profiles'][profile]['clauses']
        texts=[v.format(subject=scope) for v in source_texts[:4]]
        if not any('{subject}' in v for v in source_texts[:4]):texts[0]='围绕'+scope+'职责，'+texts[0]
        if mode=='辅助或初级':texts=['在指导下，'+v for v in texts]
        elif mode=='管理':texts.insert(0,'在授权范围内，组织'+scope+'相关工作，分配任务、协调资源并跟进质量与进度。')
        elif mode=='资深专业':texts.append('分析'+scope+'中的复杂专业问题，复核关键成果并提供方法指导，提出标准和流程改进建议。')
    clauses=[{'number':n,'text':text,'basis_type':'v2层级推定补写','source_refs':[],
      'basis_note':'依据v1同系列岗位、职业梯队与专业职责框架推定；不冒充原表新增岗位或逐字引文。'} for n,text in enumerate(texts,1)]
    return '\n'.join(f"{c['number']}、{c['text']}" for c in clauses),clauses,profile,mode,near['id']


def plan():
    master=read(BASE);policy=read(POLICY);rules=read(JD_RULES);by={j['id']:j for j in master['jobs']}
    issues={r['job_id'] for r in read(ROOT/policy['known_jd_issues_path'])};taken=collections.defaultdict(list)
    def register(j):
        for label in {j['name'],*j.get('aliases',[])}:
            for key in {norm(label),norm(QUAL.sub('',label))}:
                if j['id'] not in {x['id'] for x in taken[key]}:taken[key].append(j)
    for j in master['jobs']:register(j)
    additions=[];slots=[];series_ledger=[]
    original_order={s['id']:n for n,s in enumerate(master['series'])}
    ordered=sorted(master['series'],key=lambda s:(-len({by[i]['level'] for i in s['job_ids']}&{3,4,5,6}),-len(s['job_ids']),original_order[s['id']]))
    for s in ordered:
        jobs=[by[i] for i in s['job_ids']];levels={j['level'] for j in jobs};missing=sorted({3,4,5,6}-levels);r=route(s,jobs,policy,issues)
        rows=[]
        for level in missing:
            name,reason=candidate(r,level,policy);status='保留专业边界';references=[];new=None
            if name:
                references=taken.get(norm(name),[])
                if references:
                    same_level=any(j['level']==level for j in references)
                    status='同名或别名已有同级岗位-跨系列核对' if same_level else '名称或别名已占用其他层级-暂不新增'
                    reason='候选名称已对应既有/本轮其他岗位，须核对跨系列覆盖；不靠加代码或层级后缀制造另一个身份。'
                else:
                    status='新增推定岗位';code=s['prefix']+str(level);newid='UJ-'+hashlib.sha256((VERSION+'|'+s['id']+'|'+str(level)).encode()).hexdigest()[:14].upper()
                    text,clauses,profile,mode,reference=make_jd(name,level,r,jobs,policy,rules)
                    new={k:copy.deepcopy(s[k]) for k in ['function_id','function_name','function_group_id','function_group_name','function_direction_id','function_direction_name']}
                    new.update(id=newid,code=code,name=name,level=level,series_id=s['id'],series_name=s['name'],series_prefix=s['prefix'],aliases=[],source_ids=[],
                      industries=sorted({i for j in jobs for i in j['industries']}),name_code_status='v2层级补全推定-待复核',identity_contains_inference=True,
                      decision='按同系列现有岗位补足3—6级工作岗位；原岗位保留',previous_job_ids=[],jd=text,jd_clauses=clauses,jd_review_status='v2专业推定补写-待复核',
                      jd_review_reasons=[reason,'新增岗位及职责不是原表逐字证实，待用户复核'],jd_profile=profile,jd_role_mode=mode,jd_version=VERSION,
                      level_completion={'version':VERSION,'basis_series_id':s['id'],'parent_job_ids':[j['id'] for j in jobs],'jd_reference_job_id':reference,'naming_type':r['type'],
                        'naming_reason':reason,'context_source_ids':sorted({z for j in jobs for z in j['source_ids']}),'sources_are_context_only':True,
                        'industry_assignment_status':'同系列范围推定，尚未同步到子库','source_level_direct_evidence':False,'requires_review':True})
                    additions.append(new);register(new)
            record={'series_id':s['id'],'series_prefix':s['prefix'],'series_name':s['name'],'function_name':s['function_name'],'group_name':s['function_group_name'],
              'level':level,'proposed_code':s['prefix']+str(level),'candidate_name':name,'naming_type':r['type'],'status':status,'reason':reason,
              'new_job_id':new['id'] if new else None,'related_jobs':[{'id':j['id'],'code':j['code'],'name':j['name'],'level':j['level'],'series_id':j['series_id']} for j in references],
              'parent_jobs':r['anchors']}
            slots.append(record);rows.append(record)
        after=levels|{r['level'] for r in rows if r['status']=='新增推定岗位'}
        series_ledger.append({'series_id':s['id'],'prefix':s['prefix'],'name':s['name'],'function_name':s['function_name'],'group_name':s['function_group_name'],
          'levels_before':sorted(levels),'levels_after':sorted(after),'missing_before':missing,'missing_after':sorted({3,4,5,6}-after),'naming_type':r['type'],
          'decision_reason':r['reason'],'added_jobs':[v['new_job_id'] for v in rows if v['new_job_id']],'slot_status_counts':dict(collections.Counter(v['status'] for v in rows)),
          'parent_jobs':r['anchors']})
    slots.sort(key=lambda r:(original_order[r['series_id']],r['level']));series_ledger.sort(key=lambda r:original_order[r['series_id']])
    additions.sort(key=lambda j:(original_order[j['series_id']],j['level']))
    return master,additions,slots,series_ledger


def write_plan():
    master,new,slots,series=plan();OUT.mkdir(exist_ok=True)
    for name,data in [('新增岗位候选.json',new),('逐层级补全审查.json',slots),('职责系列层级审查.json',series)]:
        (OUT/name).write_bytes(encode(data))
    stats={'existing_jobs':len(master['jobs']),'series':len(series),'gap_slots_before':len(slots),'added_jobs':len(new),
      'slot_status_counts':dict(collections.Counter(r['status'] for r in slots)),'added_by_type':dict(collections.Counter(j['level_completion']['naming_type'] for j in new)),
      'complete_series_before':sum(not s['missing_before'] for s in series),'complete_series_after':sum(not s['missing_after'] for s in series)}
    print(json.dumps(stats,ensure_ascii=False,indent=2))


def rows_all(master,new_ids):
    header=['岗位ID','岗位代码','岗位名称','层级','规范职能','浏览分类','专业方向','职责系列','本版属性','适用行业','岗位JD','JD专业框架','职责角色','主库版本','来源状态']
    return [header]+[[j['id'],j['code'],j['name'],str(j['level']),j['function_name'],j['function_group_name'],j['function_direction_name'],j['series_name'],
      'v2新增推定-待复核' if j['id'] in new_ids else 'v1原岗-原样保留','、'.join(j['industries']),j['jd'],j['jd_profile'],j['jd_role_mode'],master['version'],
      '同系列上下文推定；没有伪造原表source_id' if j['id'] in new_ids else 'v1原有来源保持'] for j in master['jobs']]


def tsv(rows):
    stream=io.StringIO(newline='');csv.writer(stream,delimiter='\t',lineterminator='\n').writerows(rows);return stream.getvalue().encode('utf-8-sig')


def materialize():
    master,new,slots,series=plan();out=copy.deepcopy(master);new_ids={j['id'] for j in new};by={j['id']:j for j in master['jobs']+new}
    backup=read(BACKUP/'备份清单.json');assert sha(BASE)==backup['master_sha256']==sha(BACKUP/'主库.json')
    for name,record in backup['files'].items():assert sha(BACKUP/name)==record['sha256'],name
    out['version']=VERSION;out['parent_version']=master['version'];out['name_code_version']=VERSION
    out['phase']='主库v2：仅新增同系列3—6级岗位；不改v1原岗，不覆盖已发布子库与HTML'
    out['quality']='新增岗位为透明标记的职业梯队推定，结构核验后供复核；原岗全部字段原样保留'
    out['v1_summary']=copy.deepcopy(master['summary']);out['v1_code_review']=copy.deepcopy(master['code_review'])
    out['previous_name_code_version']=master['name_code_version'];out['previous_code_review']=copy.deepcopy(master['code_review'])
    out['v1_classification_review']=copy.deepcopy(master.get('classification_review'))
    out['classification_review']={'version':VERSION,'policy':'仅为原系列追加新岗，分类ID及名称保持；字典人数按v2实际岗位重算。'}
    out['code_review']={'version':VERSION,'status':'ADDITIVE_LEVEL_COMPLETION_FOR_REVIEW','five_characters':True,'existing_job_fields_changed':0,
      'new_code_policy':'同一系列原四位前缀+新增实际层级3/4/5/6；不迁移或重新分配原码。','subcatalogs_updated':False,'html_updated':False}
    out['level_completion_review']={'version':VERSION,'policy_path':POLICY.relative_to(ROOT).as_posix(),'original_master_sha256':sha(BASE),
      'backup_path':BACKUP.relative_to(ROOT).as_posix(),'new_roles_inferred':True,'new_original_source_ids_fabricated':False,
      'industry_membership_policy':'新增岗位取同系列行业并集作为建议范围；待v2确认后另行迁移子库。'}
    out['jobs'].extend(new)
    for s in out['series']:s['job_ids']+= [j['id'] for j in new if j['series_id']==s['id']]
    for g in out['function_groups']:
        js=[j for j in out['jobs'] if j['function_group_id']==g['id']]
        g.update(jobs=len(js),professional_directions=len({j['function_direction_id'] for j in js}),series=len({j['series_id'] for j in js}),code_blocks=dict(sorted(collections.Counter(j['code'][:2] for j in js).items())))
    for d in out['function_directions']:
        js=[j for j in out['jobs'] if j['function_direction_id']==d['id']];d.update(jobs=len(js),series=len({j['series_id'] for j in js}))
    added_by_level=dict(sorted(collections.Counter(str(j['level']) for j in new).items()))
    statuses=dict(collections.Counter(r['status'] for r in slots));initial=sum(j['level'] in {3,4,5,6} for j in master['jobs'])
    summary={'jobs':len(out['jobs']),'existing_jobs_retained':len(master['jobs']),'added_jobs':len(new),'functions':len({j['function_id'] for j in out['jobs']}),
      'function_groups':len(out['function_groups']),'professional_directions':len(out['function_directions']),'series':len(out['series']),
      'target_levels':[3,4,5,6],'target_slots':len(series)*4,'filled_target_slots_before':initial,'filled_target_slots_after':initial+len(new),
      'gap_slots_before':len(slots),'gap_slots_after':len(slots)-len(new),'complete_series_before':sum(not s['missing_before'] for s in series),
      'complete_series_after':sum(not s['missing_after'] for s in series),'added_by_level':added_by_level,'gap_decisions':statuses,
      'new_clause_counts':dict(sorted(collections.Counter(str(len(j['jd_clauses'])) for j in new).items())),
      'existing_names_changed':0,'existing_codes_changed':0,'existing_jds_changed':0,'new_series_or_letter_blocks':0,
      'identity_inferences':sum(j['identity_contains_inference'] for j in out['jobs']),'paused_original_sources_retained':23,'html_updated':False,'subcatalogs_updated':False}
    out['summary']=summary
    assert out['jobs'][:len(master['jobs'])]==master['jobs']
    assert len(out['jobs'])==len({j['id'] for j in out['jobs']})==len({j['code'] for j in out['jobs']})==len({norm(j['name']) for j in out['jobs']})
    assert all(j['code']==j['series_prefix']+str(j['level']) and re.fullmatch('[A-Z]{2}[0-9]{2}[3-6]',j['code']) for j in new)
    assert all(4<=len(j['jd_clauses'])<=5 and j['source_ids']==[] and j['identity_contains_inference'] for j in new)
    inputs=[BASE,POLICY,JD_RULES,BACKUP/'备份清单.json',ROOT/read(POLICY)['known_jd_issues_path'],Path(__file__)]
    report={'version':VERSION,'status':'STRUCTURE_VERIFIED_ADDITIVE_WORKING_MASTER','parent_version':master['version'],'all_passed':True,'summary':summary,
      'comparison_scope':'逐个审查原3077个职责系列、4104个3—6级空位；原9111岗全部字段逐值保留，新岗按本系列格式推定并单列依据。',
      'input_sha256':{p.relative_to(ROOT).as_posix():sha(p) for p in inputs},
      'limits':'同名/别名占用仅提示跨系列覆盖或层级冲突，不能当作专业身份已裁定；专业边界保留空位，不机械补齐。新增岗没有原表逐字岗位证据，全部标记待复核。'}
    allrows=rows_all(out,new_ids);newrows=[allrows[0]]+[r for r in allrows[1:] if r[0] in new_ids]
    files={'主库v2.json':encode(out),'新增岗位候选.json':encode(new),'逐层级补全审查.json':encode(slots),'职责系列层级审查.json':encode(series),
      '核验报告.json':encode(report),'主库v2岗位名称代码及JD.tsv':tsv(allrows),'仅新增岗位名称代码及JD.tsv':tsv(newrows)}
    return files,out,new,slots,series,allrows,newrows


def export(out,new,slots,series,allrows,newrows):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    workbook=Workbook();workbook.remove(workbook.active)
    def sheet(title,rows):
        ws=workbook.create_sheet(title)
        for index,row in enumerate(rows,1):
            ws.append(row)
            for col,value in enumerate(row,1):
                if isinstance(value,str):ws.cell(index,col).data_type='s'
        ws.freeze_panes='D2';ws.auto_filter.ref=ws.dimensions
        for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='254B70')
        for n,label in enumerate(rows[0],1):ws.column_dimensions[ws.cell(1,n).column_letter].width=90 if 'JD' in label else 48 if '说明' in label or '原因' in label or '参考' in label else 24
        for row in ws.iter_rows():
            for c in row:c.alignment=Alignment(vertical='top',wrap_text=True)
    sheet('主库v2全部岗位',allrows);sheet('仅新增岗位',newrows)
    gaprows=[['系列ID','系列前缀','系列名称','规范职能','浏览分类','缺失层级','拟用代码','候选名称','处置','原因','相关已占用岗位','命名类型']]+[[r['series_id'],r['series_prefix'],r['series_name'],r['function_name'],r['group_name'],str(r['level']),r['proposed_code'],r['candidate_name'] or '',r['status'],r['reason'],'；'.join(f"{j['code']} {j['name']}（{j['level']}级）" for j in r['related_jobs']),r['naming_type']] for r in slots]
    sheet('3至6级缺口审查',gaprows);sheet('保留边界及名称占用',[gaprows[0]]+[r for r in gaprows[1:] if r[8]!='新增推定岗位'])
    sheet('全部职责系列',[['系列ID','前缀','系列名称','规范职能','浏览分类','原层级','v2层级','原缺失3至6级','v2仍缺层级','新增岗数','命名类型','审查说明']]+[[s['series_id'],s['prefix'],s['name'],s['function_name'],s['group_name'],','.join(map(str,s['levels_before'])),','.join(map(str,s['levels_after'])),','.join(map(str,s['missing_before'])),','.join(map(str,s['missing_after'])),len(s['added_jobs']),s['naming_type'],s['decision_reason']] for s in series])
    new_ids={x['id'] for x in new};old_counts=collections.Counter(j['function_name'] for j in out['jobs'] if j['id'] not in new_ids);new_counts=collections.Counter(j['function_name'] for j in new)
    sheet('263职能新增统计',[['规范职能','原岗数','新增岗数','v2总数']]+[[f,old_counts[f],new_counts[f],old_counts[f]+new_counts[f]] for f in sorted(old_counts)])
    sheet('规则与说明',[['项目','说明'],['版本',out['version']],*[[k,json.dumps(v,ensure_ascii=False)] for k,v in out['summary'].items()],*[[f'原则{n}',p] for n,p in enumerate(read(POLICY)['principles'],1)],['原主库及下游','v1已按字节备份；现有子库与HTML不回写。v2新增为推定，不伪造原表来源。']])
    workbook.save(OUT/'BMJD主库v2_层级补全.xlsx')


def main():
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--plan',action='store_true');p.add_argument('--check',action='store_true');a=p.parse_args()
    if a.plan:write_plan();return
    files,out,new,slots,series,allrows,newrows=materialize()
    if a.check:
        for n,b in files.items():assert (OUT/n).read_bytes()==b,n
    else:
        for n,b in files.items():(OUT/n).write_bytes(b)
        export(out,new,slots,series,allrows,newrows)
    print(json.dumps(out['summary'],ensure_ascii=False,indent=2))


if __name__=='__main__':main()
