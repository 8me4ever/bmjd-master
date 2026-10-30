"""Export the reviewed JD catalog to Excel (requires openpyxl).

The writer and independent audit otherwise use only the Python standard library.
Existing exports are checked and never silently overwritten.
"""
import argparse
import collections
import shutil
from datetime import datetime
from pathlib import Path

from audit_job_codes import spreadsheet_rows
from audit_job_jds import DIRECTORY, compare_export, read, rows_for_jobs


def export(download=None):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    master = read(DIRECTORY/'岗位JD主库.json')
    report = read(DIRECTORY/'核验报告.json')
    ledger = read(DIRECTORY/'逐岗JD审查台账.json')
    jobs = master['jobs']
    expected = rows_for_jobs(jobs)
    path = DIRECTORY/'全部岗位名称与JD.xlsx'
    if path.exists():
        if not compare_export(expected, spreadsheet_rows(path, '岗位名称与JD')):
            raise ValueError('现有Excel内容不同，不覆盖用户已有工作。')
    else:
        workbook = Workbook()
        workbook.properties.creator = 'Codex'
        workbook.properties.title = 'BMJD 全部岗位名称与JD'
        workbook.properties.created = datetime(2026, 10, 29)
        workbook.properties.modified = datetime(2026, 10, 29)
        main = workbook.active
        main.title = '岗位名称与JD'
        for row in expected:
            main.append(row)
        main.freeze_panes = 'D2'
        for index, width in enumerate([28, 13, 42, 24, 35, 8, 14, 64, 115], 1):
            main.column_dimensions[get_column_letter(index)].width = width
        for index, job in enumerate(jobs, 2):
            main.row_dimensions[index].height = max(75, len(job['jd_clauses'])*23)
            if job['jd_review_status'] == '重写':
                main.cell(index, 7).fill = PatternFill('solid', fgColor='FFF2CC')
        changes = workbook.create_sheet('重写前后对照')
        changes.append(['岗位ID', '岗位代码', '岗位名称', '职能', '重写原因', '旧稿名称', '旧稿JD', '新JD'])
        for row in ledger:
            if row['status'] != '重写':
                continue
            candidates = row['legacy_candidates']
            changes.append([row['job_id'], row['code'], row['name'], row['function_name'], '；'.join(row['reasons']),
                            '\n'.join(c['old_name'] for c in candidates),
                            '\n\n'.join(c['old_name']+'\n'+c['old_jd'] for c in candidates), row['new_jd']])
        for column, width in zip('ABCDEFGH', [28, 13, 42, 24, 64, 42, 110, 110]):
            changes.column_dimensions[column].width = width
        for index in range(2, changes.max_row+1):
            changes.row_dimensions[index].height = 145
        stats = workbook.create_sheet('职能统计')
        stats.append(['职能', '岗位数', '匹配保留', '重写'])
        counts = collections.defaultdict(collections.Counter)
        for job in jobs:
            counts[job['function_name']][job['jd_review_status']] += 1
        for function, values in sorted(counts.items()):
            stats.append([function, sum(values.values()), values['匹配保留'], values['重写']])
        stats.column_dimensions['A'].width = 35
        for column in 'BCD':
            stats.column_dimensions[column].width = 18
        notes = workbook.create_sheet('使用说明')
        summary = report['summary']
        for row in [
            ['项目', '说明'],
            ['版本', master['version']],
            ['范围', f"9,111岗；匹配保留{summary['retained']}岗，重写{summary['rewritten']}岗。"],
            ['新稿格式', '沿用1、2、3、编号；每份新稿4或5条；不新增资格、学历或汇报线。'],
            ['审查方法', 'Codex按当前名称、职能、专业职责及来源审查；显式规则覆盖全量，重点边界案例另作自审。'],
            ['低词语覆盖', '仅作为排查线索，不单独触发重写。'],
            ['保留原则', '匹配稿保留完整正文和逐条原文引用；4—5条限制仅用于新稿。'],
            ['重写原因', '包含职责错配、空泛、原身份拆分或归并、证据不足、权限不符及通用化；原因可重叠。'],
            ['写作口径', '尽量去除特定企业和行业场景，保留医药、工程、教育等岗位必要专业对象。'],
            ['补写依据', '新写条款均标专业补写；旧稿及真实原文另存，不把新稿伪称原文摘录。'],
            ['身份与代码', '本轮不改岗位ID、代码、名称、职能、系列和层级；9,111个代码仍唯一。'],
            ['推定保留', '3,489项身份工作假设继续保留；JD补写不能反向证明身份假设。'],
            ['职责边界', '主任医师等专业职称不自动取得团队人事管理权限；初级及辅助稿注明在指导下。'],
            ['追溯', '重写前后对照见第二页；完整逐岗来源、原稿、引文及规则见仓库JSON台账。'],
        ]:
            notes.append(row)
        notes.column_dimensions['A'].width = 23
        notes.column_dimensions['B'].width = 115
        for sheet in workbook:
            sheet.freeze_panes = sheet.freeze_panes or 'A2'
            sheet.auto_filter.ref = sheet.dimensions
            sheet.row_dimensions[1].height = 28
            for cell in sheet[1]:
                cell.fill = PatternFill('solid', fgColor='17365D')
                cell.font = Font(bold=True, color='FFFFFF')
            for row in sheet.iter_rows():
                for cell in row:
                    cell.alignment = Alignment(vertical='top', wrap_text=True)
                    if isinstance(cell.value, str):
                        if len(cell.value) > 32767:
                            raise ValueError('Excel单元格超限，必须拆分后再导出。')
                        cell.data_type = 's'
        workbook.save(path)
        assert compare_export(expected, spreadsheet_rows(path, '岗位名称与JD'))
    if download:
        download = Path(download)
        download.parent.mkdir(parents=True, exist_ok=True)
        if download.exists() and download.read_bytes() != path.read_bytes():
            raise ValueError('下载副本已有不同内容，不覆盖。')
        if not download.exists():
            shutil.copyfile(path, download)
    print(f'Excel已核对：{len(jobs)}岗；{path}')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--download', type=Path)
    args = parser.parse_args()
    export(args.download)
