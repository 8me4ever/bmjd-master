"""Independent artifact audit of original-class BMJD recoding; no allocator import."""
import collections
import csv
import hashlib
import json
import re
import unicodedata
from pathlib import Path
from audit_unified_master_codes import xlsx

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'原版分类编码收口_20261030'
def read(p):return json.loads(p.read_text('utf-8'))
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def canon(c):return unicodedata.normalize('NFKC',str(c)).strip().upper()


def audit():
    old=read(ROOT/'主库职能序列整合_20261030/主库.json');new=read(OUT/'主库.json')
    metadata=read(ROOT/'全量名称代码统一_20261029/全量原表元数据.json')
    src={r['source_id']:{**r['原版BMJD'],'code':canon(r['原版BMJD']['code'])} for r in metadata}
    maps=read(OUT/'逐岗原版代码对照.json');ledger=read(OUT/'系列原分类与前缀依据.json')
    report=read(OUT/'收口核验报告.json');summary=report['summary'];ext=summary['new_blocks']
    by={j['id']:j for j in new['jobs']};before={j['id']:j for j in old['jobs']}
    series={s['id']:s for s in new['series']};oldseries={s['id']:s for s in old['series']}
    protected={'GM001','HR963','DP203','DP204','DP205','DP206','AC435','CT205'}
    core={'人力资源':'HR','财务':'FA','综合管理':'GM','行政':'OA','信息技术':'IT','市场':'MK','销售':'MS','销售管理':'SA','法律合规':'LE','质量管理':'QA','EHS':'SE','物流管理':'LO','物流':'LO','仓储':'LO','运输':'LO','进出口事务':'LO'}
    valid=lambda c:bool(re.fullmatch(r'[A-Z]{2}[0-9]{2}[1-7]',c))
    checks={
      '9111岗位名称ID及五位码唯一':len(by)==len(new['jobs'])==len({j['name'] for j in by.values()})==len({j['code'] for j in by.values()})==9111,
      '全部代码严格五位且层级末位不变':all(valid(j['code']) and j['code'][-1]==str(j['level'])==str(before[j['id']]['level']) for j in by.values()),
      '不含大小写全半角及空白码冲突':len({canon(j['code']) for j in by.values()})==9111 and all(canon(j['code'])==j['code'] for j in by.values()),
      '3077职责系列前四位唯一':len(series)==len({s['prefix'] for s in series.values()})==3077,
      '系列ID及岗位ID顺序保持':list(series)==list(oldseries) and list(by)==list(before),
      '名称JD层级分类行业来源等非编码岗位字段逐值保持':all({k:v for k,v in j.items() if k not in {'code','series_prefix'}}=={k:v for k,v in before[i].items() if k not in {'code','series_prefix'}} for i,j in by.items()),
      '系列除prefix外字段逐值保持':all({k:v for k,v in s.items() if k!='prefix'}=={k:v for k,v in oldseries[i].items() if k!='prefix'} for i,s in series.items()),
      '原263规范职能及95展示导航不作重分类':len({j['function_id'] for j in by.values()})==263 and new['function_groups']==old['function_groups'] and new['function_directions']==old['function_directions'],
      '岗位码前四位与职责系列一一对应':all(j['series_prefix']==j['code'][:4]==series[j['series_id']]['prefix'] and j['id'] in series[j['series_id']]['job_ids'] for j in by.values()),
      '16类明确基础职能主码全部满足':all(j['code'][:2]==core[j['function_name']] for j in by.values() if j['function_name'] in core),
      '人力资源60岗HR财务109岗FA综合60岗GM':all(sum(j['function_name']==f for j in by.values())==n and all(j['code'].startswith(b) for j in by.values() if j['function_name']==f) for f,n,b in [('人力资源',60,'HR'),('财务',109,'FA'),('综合管理',60,'GM')]),
      '产品108岗按用户要求全部PP':sum(j['function_group_name']=='产品' for j in by.values())==108 and all(j['code'].startswith('PP') for j in by.values() if j['function_group_name']=='产品'),
      '既有保护码全部保持原身份':all(by[i]['code']==j['code'] for i,j in before.items() if j['code'] in protected),
      '19226有效来源与23暂缓不混淆':len(metadata)==len(src)==19226 and summary['paused_source_rows']==23 and all(z in src for j in by.values() for z in j['source_ids']),
      '主版本名码版本及编码核验版本一致':new['version']==new['name_code_version']==new['code_review']['version']==report['version'],
      '新段仅PD_PE_RD各一个且同首字母':set(ext)=={'PD','PE','RD'} and len(set(ext.values()))==3 and all(k[0]==v[0] for k,v in ext.items()),
    }
    used=set()
    for p in report['input_sha256']:
        if p.endswith('主库.json') or p.endswith('统一主库.json'):
            d=read(ROOT/p)
            if isinstance(d,dict) and 'jobs' in d:used.update(j['code'][:2] for j in d['jobs'])
        elif p.endswith('名称代码主库.json'):
            used.update(j['code'][:2] for j in read(ROOT/p)['jobs'])
    for r in metadata:
        for key in ['原版BMJD','最新输入_修改版6']:used.add(canon(r[key]['code'])[:2])
    checks['3个新字母段未占用原始两版和7个工作版']=not(set(ext.values())&used)
    from openpyxl import load_workbook
    raw_used=set();raw_files=[]
    for folder in ['原始资料/原版BMJD','原始资料/最新输入_修改版6']:
        for p in sorted((ROOT/folder).glob('*.xlsx')):
            wb=load_workbook(p,read_only=True,data_only=True)
            for value, in wb.worksheets[0].iter_rows(min_row=2,min_col=2,max_col=2,values_only=True):
                c=canon(value)
                if valid(c):raw_used.add(c[:2])
            wb.close();raw_files.append(p.relative_to(ROOT).as_posix())
    checks['原始22个Excel含暂缓行直接复读确认新字母段未占用']=len(raw_files)==22 and not(set(ext.values())&raw_used)
    errors=[];forced=collections.Counter()
    for r in ledger:
        s=series[r['series_id']];js=[by[i] for i in s['job_ids']]
        entries=[src[z] for j in js for z in j['source_ids']]
        own={e['code'][:2] for e in entries if valid(e['code'])}
        allowed={'PP'} if js[0]['function_group_name']=='产品' else {core[s['function_name']]} if s['function_name'] in core else set(own)
        if any(before[j['id']]['code'] in protected for j in js):allowed.add(oldseries[s['id']]['prefix'][:2])
        if len(allowed)==1:forced[next(iter(allowed))]+=1
        p=s['prefix'][:2]
        ok=(p in allowed or any(p==n and b in allowed for b,n in ext.items()))
        if not (ok and set(r['own_original_blocks'])==own and set(r['allowed_original_blocks'])==allowed and r['source_ids']==sorted({z for j in js for z in j['source_ids']})):errors.append(s['id'])
    checks['全部系列有自身原版字母或明确规则新扩展依据']=not errors
    checks['三个原段各自强制系列超100证明至少需3个新段']=all(forced[b]>100 for b in ext) and dict(sorted(forced.items()))==summary['forced_original_series_counts']
    checks['3077系列台账无遗漏无重复']=len(ledger)==len({r['series_id'] for r in ledger})==3077 and {r['series_id'] for r in ledger}==set(series)
    checks['9111逐岗映射无遗漏无重复']=len(maps)==len({r['job_id'] for r in maps})==9111 and {r['job_id'] for r in maps}==set(by)
    map_errors=[]
    for r in maps:
        j=by[r['job_id']];originals={src[z]['code'] for z in j['source_ids'] if valid(src[z]['code'])}
        if not (r['current_code']==j['code'] and r['previous_code']==before[j['id']]['code'] and set(r['original_codes'])==originals and r['original_match']==(j['code'] in originals) and r['changed']==(j['code']!=before[j['id']]['code'])):map_errors.append(j['id'])
    checks['原五位码保留计数逐岗可复算']=not map_errors and sum(r['original_match'] for r in maps)==summary['current_original_code_matches']
    checks['原两位系列数字保留数逐岗可复算']=all(r['original_series_digits_match']==any(r['current_code'][2:4]==c[2:4] for c in r['original_codes']) for r in maps) and sum(r['original_series_digits_match'] for r in maps)==summary['original_series_digits_matches']
    checks['对照上一版改码数可复算']=sum(r['changed'] for r in maps)==summary['changed_from_previous']
    checks['扩展系列及岗位数可复算']=sum(bool(r['extension_of']) for r in ledger)==summary['new_extension_series'] and sum(bool(r['extension_of']) for r in maps)==summary['new_extension_jobs']
    checks['比前版保留更多原版五位码']=summary['current_original_code_matches']>summary['previous_original_code_matches']
    mapping_by={r['job_id']:r for r in maps}
    with (OUT/'岗位名称代码及JD.tsv').open(encoding='utf-8-sig',newline='') as f:tsv=list(csv.reader(f,delimiter='\t'))
    checks['TSV全9111岗位ID代码名称JD顺序一致']=len(tsv)==9112 and all(r[0]==j['id'] and r[1]==j['name'] and r[8]==j['code'] and r[-1]==j['jd'] and r[9]==','.join(mapping_by[j['id']]['original_codes']) for r,j in zip(tsv[1:],new['jobs']))
    book=xlsx(OUT/'原版分类岗位代码及JD.xlsx')
    checks['实际Excel全量字段与TSV逐值一致']=book['全量岗位及原分类']==tsv
    checks['实际Excel改码及扩展页完整']=len(book['本轮改码对照'])-1==summary['changed_from_previous'] and len(book['字母族新增段'])-1==summary['new_extension_jobs']
    owner={j['code']:j for j in before.values()};reuses=[{'code':j['code'],'previous_job_id':owner[j['code']]['id'],'previous_name':owner[j['code']]['name'],'current_job_id':j['id'],'current_name':j['name']} for j in by.values() if j['code'] in owner and owner[j['code']]['id']!=j['id']]
    checks['跨版本同码异身份披露准确']=read(OUT/'跨版本同码身份变化.json')==reuses and len(reuses)==summary['cross_version_code_reassignments']
    checks['旧原始文件主库HTML子库及新增工具指纹不变']=all(sha(ROOT/p)==h for p,h in report['input_sha256'].items())
    manifest=read(ROOT/'项目文件清单.json')['files']
    protected_files={p:d['sha256'] for p,d in manifest.items() if p.startswith(('原始资料/','行业一级子库_','行业专业JD优化_','BMJD持续补全工具_','completion_tool/','BMJD主库复核页面_','主库职能序列整合_','BMJD客户查岗平台_','customer_portal/'))}
    checks['原始Excel及全部下游文件按旧清单逐个未变']=all(sha(ROOT/p)==h for p,h in protected_files.items())
    checks['标明主库先行下游未迁移']=not new['code_review']['subcatalogs_updated'] and not new['code_review']['html_updated']
    assert all(checks.values()),{k:v for k,v in checks.items() if not v}
    result={'version':new['version'],'all_passed':True,'passed':sum(checks.values()),'total':len(checks),'checks':checks,'summary':summary,
      'scope':'产物结构、原分类证据及数值对照的独立核验；优化目标由有界整数代价的联合最小权匹配求解，不将原始前缀存在等同全部英文展开获证。',
      'evidence_errors':errors,'mapping_errors':map_errors,'sha256':{p.relative_to(ROOT).as_posix():sha(p) for p in sorted(OUT.iterdir()) if p.suffix in {'.json','.tsv','.xlsx'} and p.name!='独立核验报告.json'},
      'raw_workbooks_checked':raw_files,'protected_source_and_downstream_sha256':protected_files}
    path=OUT/'独立核验报告.json';path.write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'passed':result['passed'],'total':result['total'],'summary':summary},ensure_ascii=False,indent=2))


if __name__=='__main__':audit()
