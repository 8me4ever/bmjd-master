"""Export the grouped master and its review tables, without changing JDs."""
import json
from pathlib import Path
from datetime import datetime
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from group_master_functions import expected_rows
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'主库职能序列整合_20261030'
def read(n):return json.loads((OUT/n).read_text('utf-8'))
def export():
    master=read('主库.json');maps=read('新旧分类与代码映射.json');rows=expected_rows(master,maps);path=OUT/'主库主序列及岗位JD.xlsx'
    if path.exists():
        w=load_workbook(path,read_only=True,data_only=False);actual=list(w['岗位主序列及JD'].values);w.close();assert actual==[tuple(r) for r in rows];return
    w=Workbook();w.remove(w.active);w.properties.creator='BMJD项目';w.properties.created=datetime(2026,10,30);w.properties.modified=datetime(2026,10,30)
    def sheet(name,rows,widths):
        s=w.create_sheet(name)
        for r in rows:
            s.append(r)
            for c in s[s.max_row]:
                if isinstance(c.value,str):c.data_type='s'
                c.alignment=Alignment(vertical='top',wrap_text=True)
        for c in s[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='244373')
        s.freeze_panes='D2' if name=='岗位主序列及JD' else 'A2';s.auto_filter.ref=s.dimensions
        for i,v in enumerate(widths,1):s.column_dimensions[s.cell(1,i).column_letter].width=v
        return s
    s=sheet('岗位主序列及JD',rows,[27,13,42,25,26,36,8,35,110,13,26])
    for n in range(2,len(rows)+1):s.row_dimensions[n].height=95
    groups=read('主序列总览.json')
    sheet('主序列总览',[['主序列','岗位数','专业方向数','职责系列数','代码段','原职能来源']]+[[g['name'],g['jobs'],g['professional_directions'],g['series'],json.dumps(g['code_blocks'],ensure_ascii=False),'、'.join(g['source_functions'])] for g in groups],[28,10,13,13,65,100])
    audit=read('分类审核台账.json')
    sheet('原职能审核对照',[['原职能','岗位数','小职能(不超过10岗)','审查结论','新主序列及岗位数','新专业方向','归并或保留依据']]+[[r['source_function_name'],r['jobs'],'是' if r['small_function'] else '否',r['decision'],json.dumps(r['target_groups'],ensure_ascii=False),'、'.join(r['direction_names']),'；'.join(r['reasons'])] for r in audit],[26,10,19,26,55,60,110])
    sheet('产品PP对照',[['岗位ID','岗位名称','原职能','主序列','专业方向','旧代码','新代码','是否改码']]+[[r['job_id'],r['name'],r['source_function_name'],r['function_group_name'],r['function_direction_name'],r['previous_code'],r['current_code'],'是' if r['code_changed'] else '否'] for r in maps if r['function_group_name']=='产品'],[27,45,25,15,30,13,13,12])
    directions=read('专业方向字典.json')
    sheet('专业方向字典',[['主序列','专业方向','原规范职能','岗位数','职责系列数']]+[[d['group_name'],d['name'],d['source_function_name'],d['jobs'],d['series']] for d in directions],[28,32,28,12,15])
    issues=read('JD边界问题.json')
    sheet('JD边界待修订',[['岗位ID','岗位代码','岗位名称','问题','本轮处置','保留的现稿JD']]+[[r['job_id'],r['code'],r['name'],r['issue'],r['action'],r['current_jd']] for r in issues],[27,13,40,80,85,110])
    sheet('说明',[['项目','说明'],['版本',master['version']],['分类层次','95个主序列、270个专业方向、3077职责系列，9111岗位原样保留；原263规范职能字段继续保留用于来源追溯。'],['产品PP','产品108岗全部PP，89岗改码；包含用户列出的六类78岗及经过职责审核的30岗相关专业。产品研发按本次用户要求PP，不再套此前该方向的R规则。'],['最小变化','非产品岗位代码不变，全部岗位名称、JD、ID、层级、行业和来源不变。不是把岗位少的专业直接删掉，也不拼造完整梯队。'],['字母段共享','原生产计划PP05仍按原专业保留，不因为使用PP就误归产品。字母段不是排他的分类主键。'],['专业方向','一级收拢导航，二级保留专业，三级仍按真实职责系列；原规范职能字段作为来源维度保留，新归属以function_group和function_direction字段为准。'],['已知JD问题','3条既有JD错配单独披露，本轮未静默改写，不将分类核验等同全部JD语义审定。'],['下游','本轮先主库与复核页，行业子库及新增工具不更新。']],[23,135])
    w.save(path);print(path.relative_to(ROOT))
if __name__=='__main__':export()
