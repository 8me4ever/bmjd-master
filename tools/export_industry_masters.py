"""Export all 11 first-level BMJD subcatalogs and a combined Excel workbook."""
import argparse
import collections
import shutil
from datetime import datetime
from pathlib import Path

from audit_job_codes import spreadsheet_rows
from split_industry_masters import DIRECTORY, INDUSTRIES, read, table_rows


def same_rows(expected, actual):
    return expected == actual


def write_sheet(workbook, title, rows, is_job=False):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    sheet = workbook.create_sheet(title)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = 'D2' if is_job else 'B2'
    sheet.auto_filter.ref = sheet.dimensions
    widths = [28, 14, 46, 24, 35, 8, 16, 53, 115] if is_job else [26]*len(rows[0])
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for cell in sheet[1]:
        cell.fill = PatternFill('solid', fgColor='17365D')
        cell.font = Font(color='FFFFFF', bold=True)
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if isinstance(cell.value, str):
                if len(cell.value) > 32767:
                    raise ValueError('Excel单元格超限，不允许截断。')
                cell.data_type = 's'
    sheet.row_dimensions[1].height = 29
    if is_job:
        for index, row in enumerate(rows[1:], 2):
            sheet.row_dimensions[index].height = max(80, (row[8].count('\n')+1)*23)
    return sheet


def new_workbook(title):
    from openpyxl import Workbook
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.creator = 'Codex'
    workbook.properties.title = title
    workbook.properties.created = datetime(2026, 10, 29)
    workbook.properties.modified = datetime(2026, 10, 29)
    return workbook


def notes(overview, industry=None):
    return [
        ['项目', '说明'],
        ['版本', overview['version']],
        ['划分口径', '按原BMJD的传媒、制造、医药、地产、教育、消费品、物流、综合服务、能源化工、金融、高科技11行业建立一级子库。'],
        ['当前子库', industry or '全部11个一级行业子库'],
        ['来源依据', '每岗以原BMJD活动来源确定所属行业，并与主库行业标签逐岗交叉核对。'],
        ['名称和代码', '所有岗位名称、ID、代码、职能、职责系列、层级和JD均沿用现有主库，不新增行业专属代码。'],
        ['跨行业共用', '共用岗位进入对应多个子库，始终保留同一名称、代码和JD；子库内无重复岗位。'],
        ['数量解释', '9,111个不同岗位，11个子库合计19,018条岗位记录；2,042个岗位跨行业共用。'],
        ['行业适用', '所属子库为本表对应的行业，全部适用行业保留主库原有全部行业标签。'],
        ['未扩配', '本轮按已有来源拆分，不向原行业没有来源的岗位自行增加适用行业。'],
        ['JD保留', '使用前一阶段审查后的JD，不进行行业化改写，不反向改名或重新编码。'],
        ['原始追溯', '子库JSON保留完整岗位原字段及行业来源ID，并附本行业原始来源总账。'],
        ['暂缓来源', '原23条暂缓来源只保留总账，不恢复为活动岗位。'],
    ]


def validate_workbook(path, sheets):
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True)
    try:
        if workbook.sheetnames != list(sheets):
            raise ValueError(f'Excel工作表范围不同：{path}')
    finally:
        workbook.close()
    for name, expected in sheets.items():
        actual = spreadsheet_rows(path, name)
        # The standard-library reader returns 9 columns; auxiliary sheets may have fewer.
        padded = [[str(v) if v is not None else '' for v in row]+['']*(9-len(row)) for row in expected]
        if not same_rows(padded, actual):
            raise ValueError(f'Excel内容不一致：{path.name}/{name}')


def export(download=None):
    overview = read(DIRECTORY/'子库总览.json')
    combined = new_workbook('BMJD 11行业一级子库')
    overview_rows = [['行业', '岗位数', '唯一代码', '职能数', '系列数', '活动来源', '暂缓来源', '跨行业共用岗位', '本行业专属岗位']]
    for item in overview['industries']:
        overview_rows.append([item['industry'], item['jobs'], item['unique_codes'], item['functions'], item['series'],
                              item['active_original_sources'], item['paused_original_sources'],
                              item['shared_across_industries'], item['jobs']-item['shared_across_industries']])
    combined_sheets = {'子库总览': overview_rows}
    write_sheet(combined, '子库总览', overview_rows)
    for industry in INDUSTRIES:
        child = read(DIRECTORY/industry/'岗位主库.json')
        rows = table_rows(child['jobs'], industry)
        stats = collections.Counter(j['function_name'] for j in child['jobs'])
        function_rows = [['职能', '岗位数']] + [[k, v] for k, v in sorted(stats.items())]
        sheets = {'岗位主库': rows, '职能统计': function_rows, '使用说明': notes(overview, industry)}
        path = DIRECTORY/industry/'岗位主库.xlsx'
        if path.exists():
            validate_workbook(path, sheets)
        else:
            workbook = new_workbook('BMJD '+industry+'一级子库')
            for name, contents in sheets.items():
                sheet = write_sheet(workbook, name, contents, name == '岗位主库')
                if name == '使用说明':
                    sheet.column_dimensions['B'].width = 115
            workbook.save(path)
            validate_workbook(path, sheets)
        combined_sheets[industry] = rows
        write_sheet(combined, industry, rows, True)
    index = read(DIRECTORY/'岗位行业索引.json')
    shared_rows = [['岗位ID', '岗位代码', '岗位名称', '职能', '职责系列', '层级', '适用行业', '共用行业数', '岗位JD']]
    from split_industry_masters import MASTER
    by_id = {j['id']: j for j in read(MASTER)['jobs']}
    for item in index:
        if len(item['industries']) > 1:
            j = by_id[item['job_id']]
            shared_rows.append([j['id'], j['code'], j['name'], j['function_name'], j['series_name'], str(j['level']),
                                '、'.join(j['industries']), len(j['industries']), j['jd']])
    combined_sheets['跨行业共用岗位'] = shared_rows
    write_sheet(combined, '跨行业共用岗位', shared_rows, True)
    combined_sheets['使用说明'] = notes(overview)
    note_sheet = write_sheet(combined, '使用说明', combined_sheets['使用说明'])
    note_sheet.column_dimensions['B'].width = 115
    path = DIRECTORY/'BMJD_11行业一级子库.xlsx'
    if path.exists():
        validate_workbook(path, combined_sheets)
    else:
        combined.save(path)
        validate_workbook(path, combined_sheets)
    if download:
        download = Path(download)
        download.parent.mkdir(parents=True, exist_ok=True)
        if download.exists() and download.read_bytes() != path.read_bytes():
            raise ValueError('下载文件已有不同内容，不覆盖。')
        if not download.exists():
            shutil.copyfile(path, download)
    print('11个行业独立Excel及总工作簿已导出并逐行核对。')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--download', type=Path)
    args = parser.parse_args()
    export(args.download)
