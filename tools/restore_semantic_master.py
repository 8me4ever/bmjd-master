"""Restore original BMJD prefix evidence; extend only authorized P/R families.

Five characters, the original level digit and distinct four-character series
prefixes remain mandatory. Existing catalogs and HTML are immutable inputs.
"""
import argparse
import collections
import copy
import csv
import hashlib
import io
import json
import re
import unicodedata
from pathlib import Path

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'原版分类编码收口_20261030'
NEW=ROOT/'主库职能序列整合_20261030/主库.json'
JD=ROOT/'岗位JD匹配审查_20261029/岗位JD主库.json'
PREV=ROOT/'主库编码统一_20261030/主库.json'
META=ROOT/'全量名称代码统一_20261029/全量原表元数据.json'
HISTORY=[ROOT/p for p in ['全库统一完成_20261029/统一主库.json','主库推定整理_20261029/统一主库.json','岗位名称代码统一_20261029/名称代码主库.json','全量名称代码统一_20261029/名称代码主库.json',str(JD.relative_to(ROOT)),str(PREV.relative_to(ROOT)),str(NEW.relative_to(ROOT))]]
CORE={'人力资源':'HR','财务':'FA','综合管理':'GM','行政':'OA','信息技术':'IT','市场':'MK','销售':'MS','销售管理':'SA','法律合规':'LE','质量管理':'QA','EHS':'SE','物流管理':'LO','物流':'LO','仓储':'LO','运输':'LO','进出口事务':'LO'}
PROTECTED={'GM001','HR963','DP203','DP204','DP205','DP206','AC435','CT205'}
VERSION='BMJD-ORIGINAL-CLASS-CODE-20261030-01'


def read(p):return json.loads(p.read_text('utf-8'))
def encode(d):return (json.dumps(d,ensure_ascii=False,indent=2)+'\n').encode()
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def canon(c):return unicodedata.normalize('NFKC',str(c)).strip().upper()


def matching(options):
    from scipy.sparse import csr_matrix
    from scipy.sparse.csgraph import min_weight_full_bipartite_matching
    slots=sorted({p for row in options for p in row});index={p:i for i,p in enumerate(slots)}
    ri=[];ci=[];weights=[]
    for i,row in enumerate(options):
        for p,w in sorted(row.items()):ri.append(i);ci.append(index[p]);weights.append(w)
    graph=csr_matrix((weights,(ri,ci)),shape=(len(options),len(slots)),dtype='float64')
    a,b=min_weight_full_bipartite_matching(graph)
    selected={int(i):slots[int(k)] for i,k in zip(a,b)}
    assert len(selected)==len(options)==len(set(selected.values()))
    return [selected[i] for i in range(len(options))]


def inventory(master,metadata):
    sources={r['source_id']:{**r['原版BMJD'],'code':canon(r['原版BMJD']['code'])} for r in metadata}
    by={j['id']:j for j in master['jobs']};rows=[]
    for s in master['series']:
        js=[by[i] for i in s['job_ids']]
        entries=[sources[z] for j in js for z in j['source_ids']]
        own={e['code'][:2] for e in entries if re.fullmatch(r'[A-Z]{2}[0-9]{2}[1-7]',e['code'])}
        originals={j['id']:{sources[z]['code'] for z in j['source_ids'] if re.fullmatch(r'[A-Z]{2}[0-9]{2}[1-7]',sources[z]['code'])} for j in js}
        if js[0]['function_group_name']=='产品':allowed={'PP'};basis='用户明确指定产品PP'
        elif s['function_name'] in CORE:allowed={CORE[s['function_name']]};basis='原版基础职能主码与用户语义原则'
        else:allowed=set(own);basis='本职责系列自身原版部门、职能和专业来源'
        assert allowed,s['id']
        pinned=any(j['code'] in PROTECTED for j in js)
        if pinned:
            allowed.add(s['prefix'][:2])
            basis+='；按既有明确保护码保留身份，例外单列'
        rows.append({'series':s,'jobs':js,'own_blocks':sorted(own),'allowed_blocks':sorted(allowed),
          'originals':originals,'basis':basis,'pinned':pinned,
          'original_functions':sorted({e['function'] for e in entries}),
          'original_systems':sorted({e['system'] for e in entries}),
          'original_specialties':sorted({e['specialty'] for e in entries}),
          'source_ids':sorted({z for j in js for z in j['source_ids']})})
    return sources,rows


def options(rows,extensions):
    all_options=[]
    for r in rows:
        s=r['series'];js=r['jobs'];allowed=set(r['allowed_blocks'])
        allowed.update(new for old,new in extensions.items() if old in allowed)
        prefixes={s['prefix']} if r['pinned'] else {b+f'{n:02}' for b in allowed for n in range(100)}
        costs={}
        for p in sorted(prefixes):
            misses=sum(p+str(j['level']) not in r['originals'][j['id']] for j in js)
            extension=sum(p[:2] not in r['allowed_blocks'] for j in js)
            numeric_misses=sum(not any(p[2:]==c[2:4] for c in r['originals'][j['id']]) for j in js)
            changed=sum(p+str(j['level'])!=j['code'] for j in js)
            costs[p]=misses*1_000_000_000_000+extension*100_000_000+numeric_misses*10_000+changed+1
        all_options.append(costs)
    # All lower-priority costs combined are less than one original-code change.
    total=sum(len(r['jobs']) for r in rows)
    assert total*100_000_000+total*10_000+total<1_000_000_000_000
    assert total*10_000+total<100_000_000
    assert total<10_000
    return all_options


def fresh_extensions(used):
    """These three original blocks demonstrably exceed 100 forced series.

    New blocks inherit the original class; the second letter is an extension
    identifier, not a fabricated claim about an official English expansion.
    """
    result={}
    for base in ['PD','PE','RD']:
        chosen=next(base[0]+c for c in 'ABCDEFGHIJKLMNOPQRSTUVWXYZ' if base[0]+c not in used and base[0]+c not in result.values())
        result[base]=chosen
    return result


def build():
    master=read(NEW);metadata=read(META);sources,rows=inventory(master,metadata)
    used=set()
    for p in HISTORY:used.update(j['code'][:2] for j in read(p)['jobs'])
    for r in metadata:
        for key in ['原版BMJD','最新输入_修改版6']:used.add(canon(r[key]['code'])[:2])
    extensions=fresh_extensions(used)
    forced=collections.Counter(r['allowed_blocks'][0] for r in rows if len(r['allowed_blocks'])==1)
    assert all(forced[b]>100 for b in extensions)
    selection=matching(options(rows,extensions))
    by_series={r['series']['id']:r for r in rows};selected={r['series']['id']:p for r,p in zip(rows,selection)}
    out=copy.deepcopy(master);out['version']=VERSION;out['parent_version']=master['version'];out['name_code_version']=VERSION
    out['phase']='原版分类与前缀依据优先、五位码不变、仅主库收口；行业库与HTML待确认后迁移'
    out['previous_code_review']=copy.deepcopy(master['code_review'])
    out['code_review']={'version':VERSION,'status':'VERIFIED_PENDING_USER_REVIEW','subcatalogs_updated':False,'html_updated':False,
      'principle':'以原版部门/职能/专业及本系列来源确定字母段；产品PP例外；字母族容量不足仅用全新P/R扩展；后两位系列数字避重，末位层级保留。',
      'navigation_groups':'95个展示大类只保留导航，不作为新造职能前缀依据。',
      'new_extensions':extensions,'english_note':'未发现全库官方英文展开词典；保留原码与原分类的关联，不对既有码进行反向拼词。'}
    for s in out['series']:s['prefix']=selected[s['id']]
    mappings=[]
    for j in out['jobs']:
        r=by_series[j['series_id']];old=j['code'];p=selected[j['series_id']]
        j['code']=p+str(j['level']);j['series_prefix']=p
        originals=r['originals'][j['id']]
        reason='保留原版五位码' if j['code'] in originals else '保留现码，已符合本系列原版字母依据' if j['code']==old else '用户指定产品PP' if j['function_group_name']=='产品' else '经授权在原字母族新开扩展，数字避重' if p[:2] in extensions.values() else '沿用原职能字母段，数字避重或保留当前层级'
        mappings.append({'job_id':j['id'],'name':j['name'],'function_name':j['function_name'],
          'group':j['function_group_name'],'direction':j['function_direction_name'],'series_id':j['series_id'],
          'series_name':j['series_name'],'level':j['level'],'previous_code':old,'current_code':j['code'],
          'original_codes':sorted(originals),'original_match':j['code'] in originals,'changed':old!=j['code'],
          'original_series_digits_match':any(p[2:]==c[2:4] for c in originals),
          'original_functions':r['original_functions'],'original_systems':r['original_systems'],
          'original_specialties':r['original_specialties'],'original_blocks':r['own_blocks'],
          'extension_of':next((b for b,n in extensions.items() if n==p[:2]),''),
          'reason':reason,'source_ids':j['source_ids']})
    ledger=[]
    for r in rows:
        s=r['series'];p=selected[s['id']]
        ledger.append({'series_id':s['id'],'name':s['name'],'function_name':s['function_name'],'jobs':len(r['jobs']),
          'previous_prefix':s['prefix'],'current_prefix':p,'own_original_blocks':r['own_blocks'],
          'allowed_original_blocks':r['allowed_blocks'],'extension_of':next((b for b,n in extensions.items() if n==p[:2]),''),
          'original_functions':r['original_functions'],'original_systems':r['original_systems'],
          'original_specialties':r['original_specialties'],'basis':r['basis'],'pinned':r['pinned'],
          'job_ids':s['job_ids'],'source_ids':r['source_ids']})
    previous_own_matches=sum(j['code'] in r['originals'][j['id']] for r in rows for j in r['jobs'])
    summary={'jobs':len(mappings),'series':len(rows),'source_rows':len(metadata),'paused_source_rows':23,
      'changed_from_previous':sum(r['changed'] for r in mappings),'retained_previous':sum(not r['changed'] for r in mappings),
      'previous_original_code_matches':previous_own_matches,'current_original_code_matches':sum(r['original_match'] for r in mappings),
      'original_code_matches_gained':sum(r['original_match'] for r in mappings)-previous_own_matches,
      'original_series_digits_matches':sum(r['original_series_digits_match'] for r in mappings),
      'original_block_or_explicit_rule_jobs':sum(not r['extension_of'] for r in mappings),
      'new_extension_jobs':sum(bool(r['extension_of']) for r in mappings),
      'new_extension_series':sum(bool(r['extension_of']) for r in ledger),
      'new_blocks':extensions,'forced_original_series_counts':dict(sorted(forced.items())),
      'changed_names':0,'changed_jds':0,'changed_levels':0,'changed_classification_fields':0,
      'unique_names':len({j['name'] for j in out['jobs']}),'unique_codes':len({j['code'] for j in out['jobs']}),
      'subcatalogs_updated':False,'html_updated':False}
    oldowner={j['code']:j for j in master['jobs']};reassigned=[]
    for j in out['jobs']:
        old=oldowner.get(j['code'])
        if old and old['id']!=j['id']:reassigned.append({'code':j['code'],'previous_job_id':old['id'],'previous_name':old['name'],'current_job_id':j['id'],'current_name':j['name']})
    summary['cross_version_code_reassignments']=len(reassigned)
    report={'version':VERSION,'status':'VERIFIED_PENDING_USER_REVIEW','summary':summary,
      'objective':'先在原分类/本系列原字母、明确基础主码、产品PP、五位码和系列唯一约束下最大化原版五位码保留数，再最少使用新扩展，再尽量保留原两位系列数字，最后最少改动上一版。不是对错误工作版求局部少改。',
      'scope':'原版19226条有效来源及23条暂缓来源；9,111岗与原版按source_ids对照，不按行或仅按名字硬配。多来源岗位匹配任一原五位码即计保留；原级与现级不同可能必然需要换码。',
      'capacity':'两位系列数字最多100系列。PD/PE/RD各自只有自身原字母依据的系列均超过100；每个超容量原段需至少一个新段，3个新段足以全量匹配。新段在原始两版及7个工作版本中从未使用。',
      'historical_exceptions':['HR96会务沿用原版历史用法，其中HR963继续执行既有保护规则，不将其作为迁走人资HR的理由。','PP05生产计划是原版历史含义；产品PP新码避开其完整码。字母段可有原版多义用法，当前完整代码唯一。','GM001、HR963、DP203—206、AC435、CT205执行既有明确保护规则。其中AC435及DP203—206并非自身原五位码，不能作为完整原码保留计数；逐系列明确标记保护例外。'],
      'limits':'原有3条JD错配以及身份推定仍保留并可回查；本轮不将原分类沿用称作每条JD均无问题。95类展示导航仍是旧归并结果，不作为新编码字典。',
      'input_sha256':{p.relative_to(ROOT).as_posix():sha(p) for p in dict.fromkeys([NEW,META,*HISTORY])}}
    header=['岗位ID','岗位名称','规范职能','展示主序列','专业方向','系列','层级','上一版代码','本版代码','原版代码集合','保留原版五位码','原版部门体系','原版职能','原版细分领域','扩展自','处置','岗位JD']
    table=[header]+[[j['id'],j['name'],j['function_name'],j['function_group_name'],j['function_direction_name'],j['series_name'],j['level'],r['previous_code'],j['code'],','.join(r['original_codes']),str(r['original_match']),'、'.join(r['original_systems']),'、'.join(r['original_functions']),'、'.join(r['original_specialties']),r['extension_of'],r['reason'],j['jd']] for j,r in zip(out['jobs'],mappings)]
    stream=io.StringIO(newline='');csv.writer(stream,delimiter='\t',lineterminator='\n').writerows(table)
    files={'主库.json':encode(out),'逐岗原版代码对照.json':encode(mappings),'系列原分类与前缀依据.json':encode(ledger),
      '收口核验报告.json':encode(report),'跨版本同码身份变化.json':encode(reassigned),'岗位名称代码及JD.tsv':stream.getvalue().encode('utf-8-sig')}
    return files,report,table,ledger


def export(table,ledger,report):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment,Font,PatternFill
    wb=Workbook();wb.remove(wb.active)
    def sheet(name,rows):
        ws=wb.create_sheet(name)
        for row_index,row in enumerate(rows,1):
            ws.append(row)
            for column_index,value in enumerate(row,1):
                if isinstance(value,str):ws.cell(row_index,column_index).data_type='s'
        ws.freeze_panes='C2';ws.auto_filter.ref=ws.dimensions
        for c in ws[1]:c.font=Font(color='FFFFFF',bold=True);c.fill=PatternFill('solid',fgColor='254B70')
        for row in ws.iter_rows():
            for c in row:c.alignment=Alignment(vertical='top',wrap_text=True)
        for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=min(60,max(16,len(str(col[0].value))*2))
    sheet('全量岗位及原分类',table)
    sheet('本轮改码对照',[table[0]]+[r for r in table[1:] if r[7]!=r[8]])
    sheet('字母族新增段',[table[0]]+[r for r in table[1:] if r[14]])
    sheet('3077系列来源',[['系列ID','系列','规范职能','岗位数','旧前缀','新前缀','本系列原字母','原部门','原职能','原专业','扩展自','依据']]+[[r['series_id'],r['name'],r['function_name'],r['jobs'],r['previous_prefix'],r['current_prefix'],','.join(r['own_original_blocks']),'、'.join(r['original_systems']),'、'.join(r['original_functions']),'、'.join(r['original_specialties']),r['extension_of'],r['basis']] for r in ledger])
    sheet('说明',[['项目','内容'],['版本',VERSION],*[[k,str(v)] for k,v in report['summary'].items()],['目标',report['objective']],['对照口径',report['scope']],['容量',report['capacity']],['历史例外','；'.join(report['historical_exceptions'])],['限制',report['limits']],['下游状态','只更新主库；已有HTML与行业库仍是前版代码，请勿混用']])
    wb.save(OUT/'原版分类岗位代码及JD.xlsx')


def main():
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--check',action='store_true');args=p.parse_args()
    files,report,table,ledger=build()
    if args.check:
        for n,b in files.items():assert (OUT/n).read_bytes()==b,n
    else:
        OUT.mkdir(exist_ok=True)
        for n,b in files.items():(OUT/n).write_bytes(b)
        export(table,ledger,report)
    print(json.dumps(report['summary'],ensure_ascii=False,indent=2))


if __name__=='__main__':main()
