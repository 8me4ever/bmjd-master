"""Export only the candidate corrected global master, never children or HTML."""
from datetime import datetime
from pathlib import Path
import json
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Alignment,Font,PatternFill
from openpyxl.utils import get_column_letter

ROOT=Path(__file__).resolve().parents[1]
DIRECTORY=ROOT/'主库编码统一_20261030'
def read(name):return json.loads((DIRECTORY/name).read_text('utf-8'))

def export():
    master=read('主库.json');mappings=read('新旧岗位代码映射.json');by_id={r['job_id']:r for r in mappings}
    header=['岗位ID','岗位代码','岗位名称','职能','职责系列','层级','适用行业','岗位JD','旧代码','本轮处置']
    rows=[header]+[[j['id'],j['code'],j['name'],j['function_name'],j['series_name'],j['level'],'、'.join(j['industries']),j['jd'],by_id[j['id']]['previous_code'],by_id[j['id']]['decision']] for j in master['jobs']]
    path=DIRECTORY/'主库名称代码及JD.xlsx'
    if path.exists():
        workbook=load_workbook(path,read_only=True,data_only=False)
        actual=list(workbook['岗位名称代码及JD'].values);workbook.close()
        assert actual==[tuple(r) for r in rows],'已有Excel内容不同，不覆盖。'
        return
    workbook=Workbook();workbook.remove(workbook.active)
    workbook.properties.title='BMJD 主库编码校正版（待确认）'
    workbook.properties.creator='BMJD项目';workbook.properties.created=datetime(2026,10,30);workbook.properties.modified=datetime(2026,10,30)
    def sheet(name,values,widths):
        ws=workbook.create_sheet(name)
        for values_row in values:
            ws.append(values_row)
            for cell in ws[ws.max_row]:
                if isinstance(cell.value,str):cell.data_type='s'
                cell.alignment=Alignment(vertical='top',wrap_text=True)
        for cell in ws[1]:cell.fill=PatternFill('solid',fgColor='244373');cell.font=Font(bold=True,color='FFFFFF')
        ws.freeze_panes='D2' if name=='岗位名称代码及JD' else 'A2';ws.auto_filter.ref=ws.dimensions
        for i,width in enumerate(widths,1):ws.column_dimensions[get_column_letter(i)].width=width
        return ws
    main=sheet('岗位名称代码及JD',rows,[27,12,42,22,35,8,35,115,12,58])
    for i in range(2,len(rows)+1):main.row_dimensions[i].height=100
    sheet('新旧码改动对照',[['岗位ID','岗位名称','职能','职责系列','层级','旧代码','新代码','处理依据']]+[[r['job_id'],r['name'],r['function_name'],r['series_name'],r['level'],r['previous_code'],r['current_code'],r['decision']] for r in mappings if r['changed']],[27,43,23,35,8,12,12,70])
    sheet('职能代码核对',[['职能','岗位数','旧字母段及数量','新字母段及数量','改码数','基础职能应使用']]+[[r['function_name'],r['jobs'],json.dumps(r['previous_blocks'],ensure_ascii=False),json.dumps(r['current_blocks'],ensure_ascii=False),r['changed_jobs'],r['core_expected_block'] or '按原版来源及职责系列核对'] for r in read('职能代码核查.json')],[27,10,45,55,12,50])
    sheet('跨版本同码身份变化',[['代码','旧岗位ID','旧岗位名称','新岗位ID','新岗位名称','回查注意']]+[[r['code'],r['previous_job_id'],r['previous_name'],r['current_job_id'],r['current_name'],r['notice']] for r in read('跨版本代码身份变更.json')],[12,27,43,27,43,85])
    report=read('核验报告.json')
    sheet('使用说明',[['项目','说明'],['状态','仅主库校正版，待用户确认；旧主库保留，子库和HTML尚未更新。'],['版本',master['version']],['数量',json.dumps(report['summary'],ensure_ascii=False)],['原版优先','保留合理既有码。按本系列原版来源恢复字母段，原码冲突时在同一有据字母段内扩展；取消两位字母职能独占。'],['唯一性','9111个标准名称与9111个代码一一对应，3077个系列前四位唯一；代码末位与层级一致。'],['历史回查','原版跨行业原码本来可能对应不同身份，校正也可能重新启用码位；旧资料必须用版本与岗位ID定位，不能无版本只按代码合并。'],['保留字段','岗位仅code和series_prefix变化。名称、ID、职能、层级、JD、行业、来源及历史判断原样保留。'],['下一步','先确认主库，再拆分行业子库，最后更新客户查询及新增HTML和配码规则。']],[24,130])
    workbook.save(path)
    print(path.relative_to(ROOT))

if __name__=='__main__':export()
